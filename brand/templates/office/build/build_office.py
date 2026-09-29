"""Build the Office versions of the business documents.

    python brand/templates/office/build/build_office.py

Writes (next to this folder):
    krabi-documents.xlsx   quotation / invoice / receipt sheets with formulas + one company-settings sheet
    quotation.docx, invoice.docx, receipt.docx, receipt-tax-invoice.docx   Word versions (totals typed by hand)

Needs openpyxl, python-docx, Pillow. Logo PNG comes from render-logo.js.
Fonts: IBM Plex Sans Thai (body) + LINE Seed Sans TH (headings); both carry Latin glyphs, so one font per run works in Office.
"""
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE.parent
LOGO = HERE / "logo-horizontal.png"

NAVY, BLUE, SKY, MIST, LINE, SLATE, INK = "041A53", "0059FF", "7FAEFF", "EEF3FC", "DCE3F0", "56627F", "0B1633"
BODY, HEAD = "IBM Plex Sans Thai", "LINE Seed Sans TH"

COMPANY = [
    ("ชื่อแบรนด์", "Krabi Digital Solutions"),
    ("ชื่อนิติบุคคล", "[ชื่อนิติบุคคลตามที่จดทะเบียน]"),
    ("ที่อยู่", "[ที่อยู่บริษัท] จ.กระบี่ [รหัสไปรษณีย์]"),
    ("เลขประจำตัวผู้เสียภาษี", "[13 หลัก]"),
    ("สาขา", "สำนักงานใหญ่"),
    ("โทร", "080-893-0617"),
    ("อีเมล", "contact@krabitech.com"),
    ("เว็บไซต์", "krabitech.com"),
    ("LINE", "@982ghtqj"),
    ("บัญชีรับเงิน", "ธนาคาร [ชื่อธนาคาร] สาขา [สาขา]\nชื่อบัญชี [ชื่อบัญชี]\nเลขที่บัญชี [xxx-x-xxxxx-x]"),
]
ITEMS = [
    ("เว็บไซต์พร้อมระบบจองห้องพัก", "ออกแบบ UX/UI 2 ภาษา (ไทย / อังกฤษ) · ระบบจองตรงพร้อมปฏิทินห้องว่าง · On-page SEO เบื้องต้น", 1, "งาน", 85000),
    ("ตั้งค่าโดเมน อีเมลบริษัท และ SSL", "", 1, "งาน", 3500),
    ("ดูแลเซิร์ฟเวอร์และสำรองข้อมูล", "อัปเดตระบบ ตรวจสอบความปลอดภัย สำรองข้อมูลรายวัน", 12, "เดือน", 2500),
]
CLIENT = ["บริษัท [ชื่อลูกค้า] จำกัด (สำนักงานใหญ่)", "[ที่อยู่ลูกค้า]", "เลขประจำตัวผู้เสียภาษี [13 หลัก]", "คุณ[ผู้ติดต่อ] · [เบอร์โทร]"]
TODAY = date.today()

DOCS = {
    "quotation": dict(
        th="ใบเสนอราคา", en="QUOTATION", no="QT", sheet="ใบเสนอราคา", client="ลูกค้า",
        meta=[("เลขที่", None), ("วันที่", "date"), ("ยืนราคาถึง", 30), ("ผู้ดูแลโปรเจกต์", "[ชื่อผู้ดูแลโปรเจกต์]")],
        notes=("เงื่อนไข", "• ชำระมัดจำ [50]% เมื่อยืนยันงาน ส่วนที่เหลือเมื่อส่งมอบ\n• รับประกันแก้ไขข้อผิดพลาดหลังส่งมอบ [3] เดือน\n• ไม่รวมค่าโดเมน โฮสติ้ง และบริการภายนอก เว้นแต่ระบุในรายการ\n• เริ่มงานภายใน [7] วันหลังได้รับมัดจำ"),
        side=("ช่องทางชำระเงิน", "bank"),
        sign=[("ผู้เสนอราคา", "Krabi Digital Solutions"), ("ผู้อนุมัติสั่งงาน", "ลูกค้า · วันที่ ____/____/______")],
    ),
    "invoice": dict(
        th="ใบแจ้งหนี้", en="INVOICE", no="INV", sheet="ใบแจ้งหนี้", client="ลูกค้า",
        meta=[("เลขที่", None), ("วันที่", "date"), ("ครบกำหนดชำระ", 7), ("อ้างอิง", "[เลขที่ QT]")],
        notes=("หมายเหตุ", "งวดที่ [1] จาก [2] · [มัดจำ 50%]\nกรุณาแจ้งหลักฐานการโอนทาง LINE @982ghtqj หรือ contact@krabitech.com"),
        side=("ช่องทางชำระเงิน", "bank"),
        sign=[("ผู้วางบิล", "Krabi Digital Solutions"), ("ผู้รับวางบิล", "วันที่ ____/____/______")],
    ),
    "receipt": dict(
        th="ใบเสร็จรับเงิน", en="RECEIPT", no="RC", sheet="ใบเสร็จรับเงิน", client="ได้รับเงินจาก",
        th_vat="ใบเสร็จรับเงิน / ใบกำกับภาษี", en_vat="RECEIPT / TAX INVOICE",
        meta=[("เลขที่", None), ("วันที่", "date"), ("อ้างอิง", "[เลขที่ INV]"), ("", "ต้นฉบับ")],
        notes=("ชำระโดย", "☐ เงินสด\n☐ โอนเงิน ธนาคาร [ชื่อธนาคาร]\n☐ เช็ค เลขที่ [เลขที่เช็ค] ลงวันที่ [วันที่]"),
        side=("หมายเหตุ", "ได้รับเงินตามรายการข้างต้นเรียบร้อยแล้ว\nใบเสร็จนี้จะสมบูรณ์เมื่อเรียกเก็บเงินตามเช็คได้แล้ว"),
        sign=[("ผู้รับเงิน", "Krabi Digital Solutions"), ("ผู้มีอำนาจลงนาม", "Krabi Digital Solutions")],
    ),
}
DOC_NO = lambda d: f"{DOCS[d]['no']}-{TODAY:%Y%m}-001"


# ============================================================ Excel
def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def xl_font(size=10, bold=False, color=INK, head=False):
    return Font(name=HEAD if head else BODY, size=size, bold=bold, color=color)


thin = Side(style="thin", color=LINE)
navy_rule = Side(style="medium", color=NAVY)
TH_DATE = '[$-th-TH,107]d mmm yyyy'
MONEY = '#,##0.00'


def build_settings(ws):
    ws.title = "ข้อมูลบริษัท"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 62
    ws["A1"] = "ข้อมูลบริษัท · แก้ที่นี่ครั้งเดียว ทุกเอกสารดึงไปใช้"
    ws["A1"].font = xl_font(14, True, NAVY, head=True)
    for i, (k, v) in enumerate(COMPANY, start=3):
        ws.cell(i, 1, k).font = xl_font(10, color=SLATE)
        c = ws.cell(i, 2, v)
        c.font = xl_font(10, True, NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.fill = fill(MIST)
        c.border = Border(bottom=thin)
    ws.row_dimensions[3 + len(COMPANY) - 1].height = 48
    ws["A15"] = "วิธีใช้"
    ws["A15"].font = xl_font(11, True, BLUE)
    tips = [
        "ข้อความใน [วงเล็บเหลี่ยม] คือช่องที่ต้องกรอก",
        "กรอกข้อมูลลูกค้า รายการ จำนวน ราคา และส่วนลด ช่องจำนวนเงินและยอดรวมคำนวณให้อัตโนมัติ",
        "VAT 7% และหัก ณ ที่จ่าย 3% เลือก ใช่ / ไม่ใช่ ในคอลัมน์ I (อยู่นอกพื้นที่พิมพ์)",
        "ใบเสร็จ + VAT = ใบเสร็จรับเงิน / ใบกำกับภาษี (บริการ: ออกใบกำกับภาษีเมื่อได้รับเงิน)",
        "ตัวอักษรบาทใช้ฟังก์ชัน BAHTTEXT ของ Excel",
        "พิมพ์: File › Print › A4 (ตั้งค่าหน้ากระดาษไว้แล้ว) · ติดตั้งฟอนต์ IBM Plex Sans Thai และ LINE Seed Sans TH ก่อน",
    ]
    for i, t in enumerate(tips, start=16):
        ws.cell(i, 1, "•  " + t).font = xl_font(10, color=INK)
    return {k: f"'ข้อมูลบริษัท'!$B${i}" for i, (k, _) in enumerate(COMPANY, start=3)}


def build_doc_sheet(wb, key, ref):
    d = DOCS[key]
    ws = wb.create_sheet(d["sheet"])
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHI", [6, 46, 10, 10, 15, 17, 3, 22, 10]):
        ws.column_dimensions[col].width = w
    inp = PatternFill()  # no tint: input cells must print clean

    # header
    img = XLImage(str(LOGO))
    img.height, img.width = 50, 50 * 1320 / 352
    ws.add_image(img, "A1")
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 20
    ws.merge_cells("D1:F1")
    ws["D1"] = d["en"] if "en_vat" not in d else f'=IF($I$38="ใช่","{d["en_vat"]}","{d["en"]}")'
    ws["D1"].font = Font(name="Consolas", size=9, bold=True, color=BLUE)
    ws["D1"].alignment = Alignment(horizontal="right", vertical="bottom")
    ws.merge_cells("C2:F3")
    ws["C2"] = d["th"] if "th_vat" not in d else f'=IF($I$38="ใช่","{d["th_vat"]}","{d["th"]}")'
    ws["C2"].font = xl_font(20 if "th_vat" not in d else 17, True, NAVY, head=True)
    ws["C2"].alignment = Alignment(horizontal="right", vertical="center")

    # company (left) + meta (right)
    rows = [f"={ref['ชื่อนิติบุคคล']}", f"={ref['ที่อยู่']}",
            f'="เลขประจำตัวผู้เสียภาษี "&{ref["เลขประจำตัวผู้เสียภาษี"]}&" · "&{ref["สาขา"]}',
            f'={ref["โทร"]}&" · "&{ref["อีเมล"]}&" · "&{ref["เว็บไซต์"]}']
    ws["A5"] = f"={ref['ชื่อแบรนด์']}"
    ws["A5"].font = xl_font(11, True, NAVY, head=True)
    ws.merge_cells("A5:B5")
    for i, f in enumerate(rows, start=6):
        ws.merge_cells(f"A{i}:B{i}")
        ws[f"A{i}"] = f
        ws[f"A{i}"].font = xl_font(9, color=SLATE if i > 7 else INK)
    for i, (label, val) in enumerate(d["meta"], start=5):
        ws.cell(i, 4, label).font = xl_font(9, color=SLATE)
        ws.merge_cells(f"E{i}:F{i}")
        c = ws.cell(i, 5)
        if val is None:
            c.value = DOC_NO(key)
            c.font = Font(name="Consolas", size=10, bold=True, color=NAVY)
        elif val == "date":
            c.value = TODAY
            c.number_format = TH_DATE
        elif isinstance(val, int):
            c.value = "=E6+%d" % val
            c.number_format = TH_DATE
        else:
            c.value = val
        if c.font.name != "Consolas":
            c.font = xl_font(10, True, NAVY)
        c.alignment = Alignment(horizontal="right")
        for col in (4, 5, 6):
            ws.cell(i, col).fill = fill(MIST)
    for col in range(1, 7):
        ws.cell(10, col).border = Border(top=navy_rule)

    # client
    ws["A11"] = d["client"]
    ws["A11"].font = xl_font(9, True, BLUE)
    for i, t in enumerate(CLIENT, start=12):
        ws.merge_cells(f"A{i}:D{i}")
        c = ws[f"A{i}"]
        c.value = t
        c.font = xl_font(11 if i == 12 else 9.5, i == 12, NAVY if i == 12 else INK, head=i == 12)
        c.fill = inp

    # items
    H = 17
    for col, t in enumerate(["#", "รายละเอียด", "จำนวน", "หน่วย", "ราคาต่อหน่วย", "จำนวนเงิน (บาท)"], start=1):
        c = ws.cell(H, col, t)
        c.font = xl_font(9, True, "FFFFFF")
        c.fill = fill(NAVY)
        c.alignment = Alignment(horizontal="center" if col == 1 else ("right" if col in (3, 5, 6) else "left"), vertical="center")
    ws.row_dimensions[H].height = 22
    first, last = H + 1, H + 15
    for n, r in enumerate(range(first, last + 1), start=1):
        ws.cell(r, 1, f'=IF(B{r}="","",{n})').font = xl_font(8, color=SLATE)
        ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="top")
        if n <= len(ITEMS):
            t, desc, q, u, p = ITEMS[n - 1]
            ws.cell(r, 2).value = CellRichText([TextBlock(InlineFont(rFont=BODY, sz=10, b=True, color=NAVY), t)] +
                                               ([TextBlock(InlineFont(rFont=BODY, sz=8.5, color=SLATE), "\n" + desc)] if desc else []))
            ws.cell(r, 3, q)
            ws.cell(r, 4, u)
            ws.cell(r, 5, p)
            ws.row_dimensions[r].height = 36 if desc else 20
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        for col in (3, 4, 5):
            ws.cell(r, col).font = xl_font(10)
            ws.cell(r, col).alignment = Alignment(horizontal="right" if col != 4 else "left", vertical="top")
        ws.cell(r, 5).number_format = MONEY
        ws.cell(r, 6, f'=IF(OR(C{r}="",E{r}=""),"",C{r}*E{r})').number_format = MONEY
        ws.cell(r, 6).font = xl_font(10, color=NAVY)
        ws.cell(r, 6).alignment = Alignment(horizontal="right", vertical="top")
        for col in range(1, 7):
            ws.cell(r, col).border = Border(bottom=thin)

    # totals
    t0 = last + 2  # 34
    lab = lambda r, t: (ws.cell(r, 5, t).__setattr__("font", xl_font(9.5, color=SLATE)), ws.cell(r, 5).__setattr__("alignment", Alignment(horizontal="right")))
    money = lambda r, f, bold=False: (ws.cell(r, 6, f).__setattr__("number_format", MONEY), ws.cell(r, 6).__setattr__("font", xl_font(10, bold, NAVY)), ws.cell(r, 6).__setattr__("alignment", Alignment(horizontal="right")))
    lab(t0, "รวมเป็นเงิน"); money(t0, f"=SUM(F{first}:F{last})", True)
    lab(t0 + 1, "ส่วนลด"); money(t0 + 1, 0); ws.cell(t0 + 1, 6).fill = inp
    lab(t0 + 2, "ยอดหลังหักส่วนลด"); money(t0 + 2, f"=F{t0}-F{t0 + 1}")
    lab(t0 + 4, "ภาษีมูลค่าเพิ่ม 7%"); money(t0 + 4, f'=IF(I{t0 + 4}="ใช่",ROUND(F{t0 + 2}*7%,2),0)')
    gr = t0 + 5
    ws.cell(gr, 5, "ยอดรวมทั้งสิ้น").font = xl_font(10, True, "FFFFFF")
    ws.cell(gr, 6, f"=F{t0 + 2}+F{t0 + 4}").font = xl_font(13, True, "FFFFFF", head=True)
    ws.cell(gr, 6).number_format = MONEY
    for col in (5, 6):
        ws.cell(gr, col).fill = fill(BLUE)
        ws.cell(gr, col).alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[gr].height = 26
    lab(gr + 1, "หัก ณ ที่จ่าย 3%"); money(gr + 1, f'=IF(I{gr + 1}="ใช่",-ROUND(F{t0 + 2}*3%,2),0)')
    lab(gr + 2, "ยอดชำระสุทธิ"); money(gr + 2, f"=F{gr}+F{gr + 1}", True)
    dv = DataValidation(type="list", formula1='"ใช่,ไม่ใช่"', allow_blank=False)
    ws.add_data_validation(dv)
    # switches live outside the print area (H:I), so they never show on the printed document
    ws.cell(t0 + 2, 8, "ตั้งค่า (ไม่พิมพ์)").font = xl_font(9, True, BLUE)
    for r, label in ((t0 + 4, "คิด VAT 7%"), (gr + 1, "หัก ณ ที่จ่าย 3%")):
        ws.cell(r, 8, label).font = xl_font(9, color=SLATE)
        c = ws.cell(r, 9, "ไม่ใช่")
        c.font = xl_font(9, True, BLUE)
        c.fill = fill("FFF8DC")
        c.alignment = Alignment(horizontal="center")
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        dv.add(c)
    assert t0 + 4 == 38, "title formula expects the VAT switch at I38"

    ws.cell(t0, 1, "จำนวนเงินตัวอักษร").font = xl_font(9, True, BLUE)
    ws.merge_cells(start_row=t0 + 1, start_column=1, end_row=t0 + 2, end_column=2)
    c = ws.cell(t0 + 1, 1, f'=BAHTTEXT(F{gr})')
    c.font = xl_font(10, True, NAVY)
    c.fill = fill(MIST)
    c.alignment = Alignment(vertical="center", indent=1)

    # notes
    n0 = gr + 4
    ws.cell(n0, 1, d["notes"][0]).font = xl_font(9, True, BLUE)
    ws.cell(n0, 4, d["side"][0]).font = xl_font(9, True, BLUE)
    ws.merge_cells(start_row=n0 + 1, start_column=1, end_row=n0 + 4, end_column=2)
    ws.merge_cells(start_row=n0 + 1, start_column=4, end_row=n0 + 4, end_column=6)
    a = ws.cell(n0 + 1, 1, d["notes"][1])
    b = ws.cell(n0 + 1, 4, f"={ref['บัญชีรับเงิน']}" if d["side"][1] == "bank" else d["side"][1])
    for c in (a, b):
        c.font = xl_font(9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.fill = inp if c is a else PatternFill()
    for r in range(n0 + 1, n0 + 5):
        ws.row_dimensions[r].height = 17

    # signatures
    s0 = n0 + 8
    for (role, name), cols in zip(d["sign"], (("A", "B"), ("D", "F"))):
        ws.merge_cells(f"{cols[0]}{s0}:{cols[1]}{s0}")
        ws.merge_cells(f"{cols[0]}{s0 + 1}:{cols[1]}{s0 + 1}")
        ws.merge_cells(f"{cols[0]}{s0 + 2}:{cols[1]}{s0 + 2}")
        for col in range(ord(cols[0]) - 64, ord(cols[1]) - 63):
            ws.cell(s0, col).border = Border(bottom=Side(style="thin", color=INK))
        ws[f"{cols[0]}{s0 + 1}"] = role
        ws[f"{cols[0]}{s0 + 1}"].font = xl_font(9.5, True, NAVY)
        ws[f"{cols[0]}{s0 + 2}"] = name
        ws[f"{cols[0]}{s0 + 2}"].font = xl_font(9, color=SLATE)
        for r in (s0 + 1, s0 + 2):
            ws[f"{cols[0]}{r}"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[s0].height = 34

    f0 = s0 + 4
    ws.merge_cells(f"A{f0}:F{f0}")
    ws[f"A{f0}"] = f'={ref["ชื่อแบรนด์"]}&" · "&{ref["เว็บไซต์"]}&" · "&{ref["โทร"]}&" · LINE "&{ref["LINE"]}'
    ws[f"A{f0}"].font = xl_font(8, color=SLATE)
    for col in range(1, 7):
        ws.cell(f0, col).border = Border(top=thin)

    # print setup: A4 portrait, one page wide
    ws.print_area = f"A1:F{f0}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.55
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.freeze_panes = None
    return ws


def build_xlsx():
    wb = Workbook()
    ref = build_settings(wb.active)
    for k in DOCS:
        build_doc_sheet(wb, k, ref)
    wb.active = 1
    path = OUT / "krabi-documents.xlsx"
    wb.save(path)
    return path


# ============================================================ Word
def rgb(h):
    return RGBColor.from_string(h)


def set_run(run, size=10, bold=False, color=INK, head=False, mono=False):
    name = "Consolas" if mono else (HEAD if head else BODY)
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(int(size * 2)))
    rpr.append(szcs)
    if bold:
        rpr.append(OxmlElement("w:bCs"))
    return run


def para(cell_or_doc, text="", size=10, bold=False, color=INK, align=None, head=False, mono=False, space_after=0):
    p = cell_or_doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    if text:
        set_run(p.add_run(text), size, bold, color, head, mono)
    return p


def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def borders(table, bottom=None, top=None, inside_h=None):
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        spec = {"bottom": bottom, "top": top, "insideH": inside_h}.get(edge)
        if spec:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(spec[0]))
            e.set(qn("w:color"), spec[1])
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tblpr.append(b)


def cell_margins(table, mm=1.6):
    tblpr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side in ("top", "bottom", "left", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(int(mm * 56.7)))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblpr.append(m)


def widths(table, mm):
    for row in table.rows:
        for c, w in zip(row.cells, mm):
            c.width = Mm(w)


def fresh(cell):
    cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
    return cell


def build_docx(key, vat=False):
    d = DOCS[key]
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(16)
    sec.top_margin, sec.bottom_margin = Mm(12), Mm(14)
    sec.footer_distance = Mm(8)
    st = doc.styles["Normal"]
    st.font.name = BODY
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:cs"), BODY)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0
    title = d["th_vat"] if vat else d["th"]

    # header: logo | title
    t = doc.add_table(rows=1, cols=2)
    borders(t, bottom=(12, NAVY))
    widths(t, [100, 78])
    l, r = t.rows[0].cells
    fresh(l).add_paragraph().add_run().add_picture(str(LOGO), height=Mm(13))
    fresh(r)
    para(r, d["en_vat"] if vat else d["en"], 8, True, BLUE, WD_ALIGN_PARAGRAPH.RIGHT, mono=True)
    para(r, title, 18 if vat else 22, True, NAVY, WD_ALIGN_PARAGRAPH.RIGHT, head=True, space_after=6)
    if key == "receipt":
        para(r, "ต้นฉบับ / สำเนา", 8.5, True, BLUE, WD_ALIGN_PARAGRAPH.RIGHT, space_after=6)
    para(doc, space_after=4)

    # parties + meta
    t = doc.add_table(rows=1, cols=3)
    borders(t)
    cell_margins(t, 1.2)
    widths(t, [62, 62, 54])
    a, b, m = (fresh(c) for c in t.rows[0].cells)
    para(a, "ผู้ออกเอกสาร", 8, True, BLUE, space_after=2)
    para(a, COMPANY[0][1], 10.5, True, NAVY, head=True)
    for k in (1, 2):
        para(a, COMPANY[k][1], 9)
    para(a, f"เลขประจำตัวผู้เสียภาษี {COMPANY[3][1]} · {COMPANY[4][1]}", 8.5, color=SLATE)
    para(a, f"{COMPANY[5][1]} · {COMPANY[6][1]}", 8.5, color=SLATE)
    para(b, d["client"], 8, True, BLUE, space_after=2)
    para(b, CLIENT[0], 10.5, True, NAVY, head=True)
    for x in CLIENT[1:]:
        para(b, x, 9, color=INK if x is CLIENT[1] else SLATE)
    shade(m, MIST)
    mt = m.add_table(rows=0, cols=2)
    for label, val in d["meta"]:
        if not label:
            continue
        row = mt.add_row().cells
        v = DOC_NO(key) if val is None else (f"{TODAY.day}/{TODAY.month}/{TODAY.year + 543}" if val == "date" else
                                             (f"{(TODAY + timedelta(days=val)).day}/{(TODAY + timedelta(days=val)).month}/{(TODAY + timedelta(days=val)).year + 543}" if isinstance(val, int) else val))
        para(fresh(row[0]), label, 8.5, color=SLATE)
        para(fresh(row[1]), v, 8.5 if val is None else 9, True, NAVY, WD_ALIGN_PARAGRAPH.RIGHT, mono=val is None)
    widths(mt, [21, 31])
    para(m, "", 4)
    para(doc, space_after=6)

    # items
    t = doc.add_table(rows=1, cols=6)
    borders(t, inside_h=(4, LINE), bottom=(4, LINE))
    cell_margins(t, 1.1)
    heads = ["#", "รายละเอียด", "จำนวน", "หน่วย", "ราคาต่อหน่วย", "จำนวนเงิน (บาท)"]
    for i, (c, h) in enumerate(zip(t.rows[0].cells, heads)):
        shade(c, NAVY)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        fresh(c)
        para(c, h, 8.5, True, "FFFFFF", WD_ALIGN_PARAGRAPH.RIGHT if i in (2, 4, 5) else (WD_ALIGN_PARAGRAPH.CENTER if i == 0 else None))
    rows = ITEMS + [("", "", "", "", "")] * 2
    sub = 0
    for n, (ti, desc, q, u, p) in enumerate(rows, start=1):
        cells = [fresh(c) for c in t.add_row().cells]
        amt = q * p if q != "" else ""
        sub += amt or 0
        para(cells[0], str(n) if ti else "", 8, color=SLATE, align=WD_ALIGN_PARAGRAPH.CENTER)
        para(cells[1], ti, 9.5, True, NAVY)
        if desc:
            para(cells[1], desc, 8.5, color=SLATE)
        para(cells[2], f"{q}" if q != "" else "", 9.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
        para(cells[3], u, 9.5)
        para(cells[4], f"{p:,.2f}" if p != "" else "", 9.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
        para(cells[5], f"{amt:,.2f}" if amt != "" else "", 9.5, color=NAVY, align=WD_ALIGN_PARAGRAPH.RIGHT)
    widths(t, [9, 83, 15, 15, 26, 30])
    para(doc, space_after=6)

    # words | totals
    t = doc.add_table(rows=1, cols=2)
    borders(t)
    widths(t, [100, 78])
    w, s = (fresh(c) for c in t.rows[0].cells)
    para(w, "จำนวนเงินตัวอักษร", 8, True, BLUE, space_after=2)
    wt = w.add_table(rows=1, cols=1)
    shade(wt.rows[0].cells[0], MIST)
    para(fresh(wt.rows[0].cells[0]), "(หนึ่งแสนสองหมื่นหกพันเจ็ดร้อยเก้าสิบห้าบาทถ้วน)" if vat else "(หนึ่งแสนหนึ่งหมื่นแปดพันห้าร้อยบาทถ้วน)", 9.5, True, NAVY)
    tt = s.add_table(rows=0, cols=2)
    v = round(sub * 0.07, 2) if vat else 0
    lines = [("รวมเป็นเงิน", f"{sub:,.2f}", False), ("ส่วนลด", "0.00", False), ("ภาษีมูลค่าเพิ่ม 7%", f"{v:,.2f}" if vat else "–", False),
             ("ยอดรวมทั้งสิ้น", f"{sub + v:,.2f}", True), ("หัก ณ ที่จ่าย 3%", "–", False), ("ยอดชำระสุทธิ", f"{sub + v:,.2f}", False)]
    for label, v, grand in lines:
        a2, b2 = (fresh(c) for c in tt.add_row().cells)
        if grand:
            shade(a2, BLUE)
            shade(b2, BLUE)
        para(a2, label, 9.5 if not grand else 10, grand, "FFFFFF" if grand else SLATE)
        para(b2, v, 10 if not grand else 13, True, "FFFFFF" if grand else NAVY, WD_ALIGN_PARAGRAPH.RIGHT, head=grand)
    para(doc, space_after=8)

    # notes
    t = doc.add_table(rows=1, cols=2)
    borders(t)
    cell_margins(t, 2.4)
    widths(t, [89, 89])
    side_text = "\n".join(COMPANY[9][1].split("\n")) if d["side"][1] == "bank" else d["side"][1]
    for c, (h, body) in zip(t.rows[0].cells, (d["notes"], (d["side"][0], side_text))):
        fresh(c)
        shade(c, "FFFFFF")
        para(c, h, 8, True, BLUE, space_after=2)
        for line in body.split("\n"):
            para(c, line, 8.8)
    para(doc, space_after=4)

    # signatures
    t = doc.add_table(rows=1, cols=2)
    borders(t)
    widths(t, [89, 89])
    for c, (role, name) in zip(t.rows[0].cells, d["sign"]):
        fresh(c)
        para(c, "", 9, space_after=14)
        para(c, "_" * 38, 9, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        para(c, role, 9.5, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER)
        para(c, name, 9, color=SLATE, align=WD_ALIGN_PARAGRAPH.CENTER)

    # footer
    fp = sec.footer.paragraphs[0]
    set_run(fp.add_run(f"{COMPANY[0][1]} · {COMPANY[7][1]} · {COMPANY[5][1]} · LINE {COMPANY[8][1]}    "), 7.5, color=SLATE)
    set_run(fp.add_run(DOC_NO(key)), 7.5, color=SLATE, mono=True)

    path = OUT / f"{key}{'-tax-invoice' if vat else ''}.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_xlsx().name)
    for k in DOCS:
        print(build_docx(k).name)
    print(build_docx("receipt", vat=True).name)
